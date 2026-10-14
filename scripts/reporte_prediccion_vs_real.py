"""
reporte_prediccion_vs_real.py
=============================
Lista por producto de la prediccion del algoritmo frente a las ventas reales.

- "Abril 2026": el mes que predice el dashboard, entrenando igual que
  scripts/entrenar.py con data/raw/Query_Result_V6.csv (que termina el
  2026-03-31); lo real sale de data/holdout/ventas_2026_04_09.csv.
- Una hoja por mes del estudio de caso (2026-04 a 2026-09) con origen rodante
  (scripts/estudio_caso.py: se entrena solo con datos anteriores a cada mes).
- "Resumen" con totales por mes, "Grupos" (suma de las predicciones por grupo de
  producto frente a lo real, con el acierto ponderado por valor de cada mes) y
  "Cómo leer" con la explicacion de columnas.

Incluye la estimacion por promedio (productos sin ventas el ultimo mes) y no
contiene datos de clientes.

Salidas:
  outputs/reports/prediccion_vs_real.xlsx
  outputs/reports/prediccion_vs_real_abril.csv   (; y UTF-8 con BOM)
  y una copia de ambos en la carpeta indicada con --copiar-a.

Requiere openpyxl (requirements-comparacion.txt).
Uso:  MPLBACKEND=Agg PYTHONIOENCODING=utf-8 .venv/Scripts/python scripts/reporte_prediccion_vs_real.py [--copiar-a RUTA]
"""
import argparse
import shutil
import sys
import tempfile
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import estudio_caso as ec  # noqa: E402
from ventas_forecast import case_study as cs  # noqa: E402
from ventas_forecast import grupos as gr  # noqa: E402
from ventas_forecast import prediccion_vs_real as pr  # noqa: E402
from ventas_forecast.model import SalesForecastModel  # noqa: E402
from ventas_forecast.paths import DATA_RAW, DEFAULT_DATASET, REPORTS, ensure_output_dirs  # noqa: E402

NOMBRES_MES = {
    "2026-04": "Abril 2026", "2026-05": "Mayo 2026", "2026-06": "Junio 2026",
    "2026-07": "Julio 2026", "2026-08": "Agosto 2026", "2026-09": "Septiembre 2026",
}
COLS_COP = ["Ventas mes anterior (COP)", "Predicción combinada (COP)",
            "Solo modelo (COP)", "Real del mes (COP)",
            "Diferencia (real − predicción)"]
ANCHOS = {"Código": 11, "Descripción": 52, "Origen": 10, "Decisión recomendada": 40,
          "Prob. crecimiento": 13, "Ventas mes anterior (COP)": 18,
          "Predicción combinada (COP)": 19, "Solo modelo (COP)": 16,
          "Real del mes (COP)": 16, "Diferencia (real − predicción)": 19,
          "Error %": 9, "Clase real": 11, "¿Acertó la decisión?": 13,
          "Probabilidad de venta": 14, "¿Se venderá? (predicción)": 14,
          "¿Se vendió?": 11, "¿Acertó venta?": 11}

COMO_LEER = [
    ("Cómo leer este archivo", None),
    ("Qué es", "Cada hoja de mes compara, producto por producto, lo que el algoritmo "
               "predijo con lo que realmente se vendió ese mes. Los montos están en pesos "
               "colombianos (COP). No hay datos de clientes."),
    ("Abril 2026", "Es el mes que predice el dashboard hoy: el modelo se entrena igual que "
                   "en scripts/entrenar.py con las ventas hasta el 31 de marzo de 2026 y "
                   "lo real sale de las ventas de abril. Es también la hoja de abril "
                   "del estudio de caso (mismo entrenamiento)."),
    ("Mayo a septiembre", "Cada mes se predice entrenando solo con datos anteriores a ese "
                          "mes (origen rodante), como si se estuviera en ese momento. Las "
                          "ventas reales de septiembre están incompletas (el registro llega "
                          "hasta el 22 de septiembre), por lo que ese mes tiende a verse "
                          "más bajo que la predicción."),
    ("Código / Descripción", "Identifican el producto."),
    ("Origen", "modelo: el producto vendió el último mes y el algoritmo (clasificador y "
               "regresor) da una decisión y un monto. promedio: el producto no vendió el "
               "último mes pero sí en los meses previos; su monto es el promedio de los "
               "últimos 3 meses y no recibe decisión del modelo."),
    ("Decisión recomendada", "Reforzar stock, Mantener stock o Reducir stock (solo origen "
                             "modelo). Para origen promedio dice \"Sin recomendación "
                             "(estimación por promedio)\"."),
    ("Prob. crecimiento", "Probabilidad que da el modelo de que el producto crezca más de "
                          "15 %. Solo origen modelo."),
    ("Ventas mes anterior (COP)", "Ventas del último mes con datos antes del mes predicho."),
    ("Predicción combinada (COP)", "Monto que muestra el sistema. Origen modelo: mezcla del "
                                   "modelo con el promedio de los últimos 3 meses (en blanco "
                                   "si el producto está fuera de los 1.900 con más ventas). "
                                   "Origen promedio: el promedio de los últimos 3 meses."),
    ("Solo modelo (COP)", "Lo que predice únicamente el Random Forest (solo origen modelo)."),
    ("Real del mes (COP)", "Lo que realmente se vendió ese mes (0 si no hubo ventas)."),
    ("Diferencia (real − predicción)", "Real menos predicción. Positivo: se vendió más de lo "
                                       "predicho; negativo: menos."),
    ("Error %", "Valor absoluto de (real − predicción) dividido entre el real. Queda en blanco "
                "si el real es 0 o no hay predicción."),
    ("Clase real", "Qué habría sido lo correcto: se compara lo vendido con las ventas del mes "
                   "anterior. Crecimiento mayor a 15 % = Reforzar; caída mayor a 10 % o "
                   "ninguna venta = Reducir; en otro caso = Mantener. En blanco para origen "
                   "promedio."),
    ("¿Acertó la decisión?", "Sí si la decisión recomendada coincide con la clase real. En "
                             "blanco para origen promedio."),
    ("Probabilidad de venta", "Probabilidad (0 a 100 %) de que el producto tenga ventas el mes "
                              "predicho, calculada por un clasificador binario entrenado solo "
                              "con datos anteriores a ese mes. Existe para todos los productos, "
                              "también los de origen promedio."),
    ("¿Se venderá? (predicción)", "Sí si la probabilidad de venta alcanza el umbral que el "
                                  "sistema fija con la validación interna del entrenamiento "
                                  "(nunca con el mes evaluado); No en otro caso."),
    ("¿Se vendió?", "Sí si el producto tuvo ventas reales ese mes (real mayor a 0); No si no "
                    "vendió."),
    ("¿Acertó venta?", "Sí si la predicción \"¿Se venderá?\" coincide con lo que pasó "
                       "(\"¿Se vendió?\")."),
    ("Grupos", "Compara, por grupo de producto (la línea o proveedor con la que se compra), "
               "la suma de las predicciones de sus productos con lo que realmente vendieron "
               "esos mismos productos en el mes. Solo cuentan los productos con monto "
               "predicho (modelo y promedio), los mismos del dashboard; los productos sin "
               "grupo van en \"Sin grupo\". Acierto % = 1 menos el valor absoluto de "
               "(real − predicción) dividido entre el real, entre 0 y 100 % (vacío si el "
               "real es 0). La parte de arriba resume cada mes y el total de los seis "
               "meses (\"Agrupado\"): el acierto ponderado es 1 menos la suma de los "
               "errores absolutos de todos los grupos dividida entre la suma de lo real, "
               "es decir, qué parte del valor vendido se acertó. Debajo está el detalle "
               "por grupo y mes. Los meses con una caída fuerte de las ventas reales "
               "(junio) bajan el acierto de todos los grupos."),
    ("Resumen", "Totales por mes: predicho por el modelo, por promedio y total, frente a lo "
                "real; cobertura frente a lo real de los 1.900 productos con más ventas; "
                "porcentaje de decisiones acertadas y error medio absoluto (MAE). "
                "\"% acierto de venta\" es el porcentaje de productos cuya venta (sí/no) "
                "se acertó entre los productos del modelo con monto y los de estimación "
                "por promedio (los que muestra el dashboard con monto); \"% acierto de "
                "venta (todos)\" incluye además los productos del modelo sin monto."),
]


def entrenar_abril(tmp):
    """Mismo procedimiento que scripts/entrenar.py (sin graficar)."""
    m = SalesForecastModel(filepath=DEFAULT_DATASET,
                           output_path=Path(tmp) / "pred_abril.csv", min_months=6)
    m.load_and_clean()
    m.build_features()
    m.cluster()
    m.prepare_split()
    m.train_classifier()
    m.train_regressor()
    m.predict_next_month()
    m.estimar_por_promedio()
    m.agregar_probabilidad_venta()
    return m


def tabla_de_modelo(model, mes):
    c = model._c
    real = cs.monthly_sales(ec.HOLDOUT, mes)
    tabla = pr.tabla_producto(model.results, real, c("id"), c("description"),
                              c("total_sold"), model.growth_threshold,
                              model.reduce_threshold)
    real_top = float(real.reindex(model._top_prods.astype(str)).fillna(0.0).sum())
    tabla["_orden"] = (tabla["Origen"] == "promedio").astype(int)
    tabla = (tabla.sort_values(["_orden", "Predicción combinada (COP)"],
                               ascending=[True, False], na_position="last")
             .drop(columns="_orden").reset_index(drop=True))
    return tabla, real_top


def _valor(v):
    return None if (isinstance(v, float) and pd.isna(v)) else v


def escribir_hoja(ws, tabla):
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    ws.append(list(tabla.columns))
    for fila in tabla.itertuples(index=False, name=None):
        ws.append([_valor(v) for v in fila])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for j, col in enumerate(tabla.columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = ANCHOS.get(col, 14)
        if col in COLS_COP:
            fmt = "#,##0"
        elif col == "Error %":
            fmt = "0.0%"
        elif col in ("Prob. crecimiento", "Probabilidad de venta"):
            fmt = "0%"
        else:
            continue
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=j).number_format = fmt


def escribir_resumen(ws, resumen):
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    df = pd.DataFrame(resumen)
    ws.append(list(df.columns))
    for fila in df.itertuples(index=False, name=None):
        ws.append([_valor(v) for v in fila])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for j, col in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = 16 if j > 1 else 18
        if "(COP)" in col:
            fmt = "#,##0"
        elif col.startswith("%") or col.startswith("Cobertura"):
            fmt = "0.0%"
        elif col.startswith("Razón"):
            fmt = "0.000"
        else:
            continue
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=j).number_format = fmt
    ws.row_dimensions[1].height = 45
    ws.cell(row=ws.max_row + 2, column=1, value=(
        "Las ventas reales de septiembre están incompletas (el registro llega hasta el "
        "22 de septiembre). Real modelo y MAE modelo usan solo los productos del modelo con "
        "monto; los productos sin monto (fuera de los 1.900 con más ventas) cuentan en "
        "las decisiones. El % de acierto de venta usa los productos del modelo con monto y "
        "los de estimación por promedio; el de \"todos\" incluye los productos del modelo "
        "sin monto."))


def escribir_grupos(ws, resumen, detalle):
    """Resumen mensual del acierto por grupo y debajo el detalle por mes y grupo."""
    from openpyxl.styles import Alignment, Font
    resumen = pd.DataFrame(resumen)
    ws.append(list(resumen.columns))
    for fila in resumen.itertuples(index=False, name=None):
        ws.append([_valor(v) for v in fila])
    n_resumen = ws.max_row
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in range(2, n_resumen + 1):
        ws.cell(row=row, column=3).number_format = "#,##0"
        ws.cell(row=row, column=4).number_format = "#,##0"
        ws.cell(row=row, column=5).number_format = "0.0%"
    for cell in ws[n_resumen]:
        cell.font = Font(bold=True)
    ws.append([])
    ws.append(["Detalle por grupo"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    cab = ["Mes"] + list(pr.COLUMNAS_GRUPO)
    ws.append(cab)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    inicio = ws.max_row + 1
    for mes, t in detalle.items():
        for fila in t.itertuples(index=False, name=None):
            ws.append([mes] + [_valor(v) for v in fila])
    for row in range(inicio, ws.max_row + 1):
        for col in (4, 5, 6):
            ws.cell(row=row, column=col).number_format = "#,##0"
        ws.cell(row=row, column=7).number_format = "0.0%"
    for letra, ancho in zip("ABCDEFG", (34, 12, 19, 19, 22, 12, 12)):
        ws.column_dimensions[letra].width = ancho


def escribir_como_leer(ws):
    from openpyxl.styles import Alignment, Font
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    for titulo, texto in COMO_LEER:
        if texto is None:
            ws.append([titulo])
            ws["A1"].font = Font(bold=True, size=14)
            continue
        ws.append([titulo, texto])
        fila = ws.max_row
        ws.cell(row=fila, column=1).font = Font(bold=True)
        ws.cell(row=fila, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=fila, column=2).alignment = Alignment(vertical="top", wrap_text=True)


def csv_abril(tabla, ruta):
    out = tabla.copy()
    for col in COLS_COP:
        out[col] = out[col].round(0).astype("Int64")
    out["Error %"] = (out["Error %"] * 100).round(1)
    out["Prob. crecimiento"] = (out["Prob. crecimiento"] * 100).round(1)
    out["Probabilidad de venta"] = (out["Probabilidad de venta"] * 100).round(1)
    out = out.rename(columns={"Error %": "Error (%)",
                              "Prob. crecimiento": "Prob. crecimiento (%)",
                              "Probabilidad de venta": "Probabilidad de venta (%)"})
    out.to_csv(ruta, sep=";", decimal=",", index=False, encoding="utf-8-sig")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--copiar-a", default=None,
                    help="carpeta donde copiar el .xlsx y el .csv")
    args = ap.parse_args()
    from openpyxl import Workbook

    ensure_output_dirs()
    tablas, resumen = {}, []
    mapa = gr.cargar_mapa(DATA_RAW / "grupos_productos.csv")
    with tempfile.TemporaryDirectory() as tmp:
        abril = entrenar_abril(tmp)
        tabla, real_top = tabla_de_modelo(abril, "2026-04")
        tablas["2026-04"] = tabla
        resumen.append(pr.resumen_mes("2026-04", tabla, real_top))
        for mes in ec.MESES[1:]:
            _, _, _, model = ec.evaluar_mes(mes, tmp, con_modelo=True)
            model.estimar_por_promedio()
            model.agregar_probabilidad_venta()
            tabla, real_top = tabla_de_modelo(model, mes)
            tablas[mes] = tabla
            resumen.append(pr.resumen_mes(mes, tabla, real_top))

    wb = Workbook()
    escribir_como_leer(wb.active)
    wb.active.title = "Cómo leer"
    escribir_resumen(wb.create_sheet("Resumen"), resumen)
    por_grupo = {m: pr.tabla_grupos_mes(t, mapa) for m, t in tablas.items()}
    resumen_grupos = pr.resumen_grupos(por_grupo)
    escribir_grupos(wb.create_sheet("Grupos"), resumen_grupos, por_grupo)
    for mes, tabla in tablas.items():
        escribir_hoja(wb.create_sheet(NOMBRES_MES[mes]), tabla)
    xlsx = REPORTS / "prediccion_vs_real.xlsx"
    wb.save(xlsx)
    csv = REPORTS / "prediccion_vs_real_abril.csv"
    csv_abril(tablas["2026-04"], csv)
    print(f"[OK] {xlsx}\n[OK] {csv}")
    if args.copiar_a:
        destino = Path(args.copiar_a)
        destino.mkdir(parents=True, exist_ok=True)
        for f in (xlsx, csv):
            shutil.copy2(f, destino / f.name)
        print(f"[OK] copiado a {destino}")
    print(pd.DataFrame(resumen).T.to_string())
    print(resumen_grupos.to_string())


if __name__ == "__main__":
    main()
